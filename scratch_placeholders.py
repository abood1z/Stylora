
import docx

def find_placeholders(file_path):
    doc = docx.Document(file_path)
    for i, para in enumerate(doc.paragraphs):
        t = para.text.lower()
        if "heading 4" in t or "heading 5" in t or "first paragraph following a heading" in t:
            print(f"[{i}] {para.text}")

if __name__ == "__main__":
    find_placeholders(r"d:\project_folder\documentationstylora.docx")
